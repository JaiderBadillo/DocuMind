import os
import zipfile
from pathlib import Path

DATASET_DIR = Path(__file__).resolve().parent.parent / "test_dataset_30_docs"

LEGAL_DIR = DATASET_DIR / "contratos_legal"
FINANCE_DIR = DATASET_DIR / "facturas_finanzas"
HR_DIR = DATASET_DIR / "talento_humano_informes"

for d in [LEGAL_DIR, FINANCE_DIR, HR_DIR]:
    d.mkdir(parents=True, exist_ok=True)

def create_pdf(filename: Path, title: str, content: str):
    """Creates a valid, compliant PDF document using reportlab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(filename), pagesize=letter)
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, 750, title[:60])
        c.setFont("Helvetica", 10)
        y = 720
        for line in content.splitlines():
            if not line.strip():
                y -= 10
                continue
            words = line.split()
            current_line = []
            for w in words:
                current_line.append(w)
                if len(" ".join(current_line)) > 80:
                    c.drawString(50, y, " ".join(current_line[:-1]))
                    y -= 14
                    current_line = [w]
                    if y < 50:
                        c.showPage()
                        c.setFont("Helvetica", 10)
                        y = 750
            if current_line:
                c.drawString(50, y, " ".join(current_line))
                y -= 14
                if y < 50:
                    c.showPage()
                    c.setFont("Helvetica", 10)
                    y = 750
        c.save()
    except Exception as e:
        print(f"Error generando PDF con reportlab: {e}")


def create_docx(filename: Path, title: str, content: str):
    """Creates a valid DOCX using python-docx if available, or lightweight zip package."""
    try:
        import docx
        doc = docx.Document()
        doc.add_heading(title, level=1)
        for p in content.split("\n\n"):
            if p.strip():
                doc.add_paragraph(p.strip())
        doc.save(str(filename))
    except Exception:
        # Minimal valid OpenXML DOCX archive fallback
        text_nodes = "".join([f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>" for line in [title] + content.splitlines() if line.strip()])
        doc_xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{text_nodes}</w:body></w:document>'
        types_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
        rels_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
        
        with zipfile.ZipFile(filename, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", types_xml)
            z.writestr("_rels/.rels", rels_xml)
            z.writestr("word/document.xml", doc_xml)

def create_txt(filename: Path, title: str, content: str):
    with open(filename, "w", encoding="utf-8") as f:
        f.write(f"{title.upper()}\n{'='*len(title)}\n\n{content}")

# ================= 30 SYNTHETIC TEST DOCUMENTS DEFINITIONS =================

def generate_all_30_documents():
    print("Generando 30 documentos de prueba sintéticos distribuidos en 3 categorías...")
    
    # ---------------- 1. CONTRATOS Y LEGAL (10 DOCUMENTOS) ----------------
    contratos = [
        ("Contrato_01_Prestacion_Servicios_Software", "Contrato de Prestación de Servicios de Desarrollo",
         "Entre los suscritos, TECNOLOGÍA ANDINA S.A.S. (Contratante) y CONSULTORÍA DIGITAL LIMITADA (Contratista), se conviene celebrar el presente contrato.\nCláusula Primera - Objeto: El Contratista se compromete a implementar la plataforma de analítica y RAG.\nCláusula Segunda - Cuantía y Valor: La cuantía total del presente acuerdo asciende a la suma de $18.500.000 COP.\nCláusula Tercera - Vigencia: El término de ejecución será de 6 meses contados a partir del acta de inicio.\nCláusula Cuarta - Penalidad: En caso de incumplimiento injustificado, la parte afectada cobrará una sanción penal pecuniaria equivalente al 20% del valor total del contrato."),
         
        ("Contrato_02_Arrendamiento_Oficinas_Comerciales", "Contrato de Arrendamiento de Bien Inmueble Comercial",
         "Comparecen INVERSIONES URBANAS S.A. (Arrendador) y SERVICIOS EMPRESARIALES DEL ORIENTE (Arrendatario).\nCláusula Primera - Inmueble: Oficina 502 del Edificio Metropolitano, Bucaramanga.\nCláusula Segunda - Canon de Arrendamiento: La suma mensual pactada es de $4.200.000 COP más cuota de administración.\nCláusula Tercera - Vigencia: Duración improrrogable de 12 meses renovable por mutuo acuerdo escrito."),
         
        ("Contrato_03_Acuerdo_Confidencialidad_NDA", "Acuerdo Mutuo de No Divulgación y Confidencialidad",
         "Las partes acuerdan proteger la información confidencial, código fuente, secretos industriales y bases de datos.\nObligación de Reserva: Ninguna de las partes podrá divulgar sin autorización previa escrita la propiedad intelectual.\nVigencia de la Reserva: El término de confidencialidad subsistirá por 5 años posteriores a la terminación."),
         
        ("Contrato_04_Mantenimiento_Servidores_Cloud", "Contrato Marco de Mantenimiento y Soporte en la Nube",
         "Contratista: CLOUD OPERATIONS COLOMBIA S.A.S. Valor del contrato: $9.600.000 COP anuales. Servicios contratados: Monitoreo 24/7 de infraestructura AWS, actualización de parches de seguridad y balanceo de carga con SLA del 99.9%."),
         
        ("Contrato_05_Cesion_Derechos_Patrimoniales", "Contrato de Cesión de Derechos Patrimoniales de Autor",
         "El autor cede de forma total y definitiva a favor de la empresa los derechos patrimoniales sobre el software DocuMind Enterprise desarrollado durante el semestre. Valor de la cesión: $5.000.000 COP."),
         
        ("Contrato_06_Licenciamiento_Software_ERP", "Acuerdo de Licenciamiento de Uso de Software Corporativo",
         "Licenciante: ENTERPRISE SYSTEMS CORP. Licenciatario: DISTRIBUIDORA DEL NORTE LTDA. Otorga 50 licencias concurrentes de usuario por valor de $22.000.000 COP con vigencia de 24 meses."),
         
        ("Contrato_07_Suministro_Equipos_Computo", "Contrato de Compraventa y Suministro de Hardware",
         "Proveedor: HARDWARE SOLUTIONS LTDA. Objeto: Suministro de 15 estaciones de trabajo Intel Core i7 con 32GB RAM. Cuantía total: $45.000.000 COP. Garantía: 3 años directa de fábrica."),
         
        ("Contrato_08_Seguro_Responsabilidad_Civil", "Póliza y Contrato de Seguro de Cumplimiento",
         "Aseguradora: SEGUROS DEL ESTADO S.A. Ampara el cumplimiento del contrato estatal de digitalización y gestión de archivos. Valor asegurado: $100.000.000 COP. Vigencia hasta diciembre de 2026."),
         
        ("Contrato_09_Prestacion_Servicios_Auditoria", "Contrato de Servicios Profesionales de Revisoría Fiscal",
         "Firma Auditora: KPMG & AUDITORES ASOCIADOS. Honorarios mensuales: $3.800.000 COP. Objeto: Dictamen de estados financieros y auditoría de cumplimiento normativo contable."),
         
        ("Contrato_10_Convenio_Pasantia_Empresarial", "Convenio Marco de Prácticas Profesionales y Pasantías",
         "Suscrito entre UNIDADES TECNOLÓGICAS DE SANTANDER (UTS) y EMPRESA INNOVADORA S.A.S. Objeto: Vinculación de estudiantes de VI semestre en proyectos de ingeniería de software con auxilio económico de $1.400.000 COP.")
    ]

    for idx, (fname, title, body) in enumerate(contratos, 1):
        if idx % 3 == 1:
            create_pdf(LEGAL_DIR / f"{fname}.pdf", title, body)
        elif idx % 3 == 2:
            create_docx(LEGAL_DIR / f"{fname}.docx", title, body)
        else:
            create_txt(LEGAL_DIR / f"{fname}.txt", title, body)

    # ---------------- 2. FACTURAS Y FINANZAS (10 DOCUMENTOS) ----------------
    facturas = [
        ("Factura_01_Servicios_Cloud_AWS", "Factura Electrónica de Venta No. FE-2026-101",
         "Emisor: AMAZON WEB SERVICES COLOMBIA S.A.S. | NIT: 901.234.567-8\nCliente: SISTEMAS INTEGRADOS UTS | NIT: 890.201.345-1\nFecha de Emisión: 15 de Marzo de 2026 | Fecha de Vencimiento: 30 de Marzo de 2026\nDetalle: Consumo instancias EC2, almacenamiento S3 y base de datos relacional.\nSubtotal: $3.200.000 COP\nIVA (19%): $608.000 COP\nTotal a Pagar: $3.808.000 COP"),
         
        ("Factura_02_Licencias_Office365", "Factura de Venta No. FE-2026-102",
         "Emisor: MICROSOFT LATAM SERVICES | NIT: 900.555.111-2\nCliente: CORPORACIÓN EDUCATIVA | NIT: 890.201.345-1\nFecha: 02 de Febrero de 2026\nConcepto: 100 suscripciones Microsoft 365 Business Standard.\nSubtotal: $6.500.000 COP\nIVA (19%): $1.235.000 COP\nTotal a Pagar: $7.735.000 COP"),
         
        ("Factura_03_Consultoria_Seguridad_Informatica", "Factura Comercial No. FC-8840",
         "Emisor: CYBERSECURITY DEFENSE LABS | NIT: 901.888.777-3\nCliente: BANCO DEL COMERCIO\nFecha: 10 de Enero de 2026\nConcepto: Test de penetración (Ethical Hacking) y auditoría ISO 27001.\nSubtotal: $12.000.000 COP\nIVA (19%): $2.280.000 COP\nTotal: $14.280.000 COP"),
         
        ("Factura_04_Equipos_Red_Cisco", "Comprobante Fiscal de Venta No. FAC-4412",
         "Emisor: REDES Y CONECTIVIDAD TOTAL LTDA.\nFecha: 22 de Abril de 2026\nDetalle: 2 Switches administrables Cisco Catalyst 48 puertos y 4 Access Points WiFi 6.\nSubtotal: $8.900.000 COP\nIVA (19%): $1.691.000 COP\nTotal: $10.591.000 COP"),
         
        ("Factura_05_Servicios_Fibra_Optica", "Factura de Telecomunicaciones No. TEL-9021",
         "Emisor: CLARO COLOMBIA EMPRESAS. Cliente: CENTRO DE INVESTIGACIÓN. Periodo: Mayo 2026. Conexión simétrica dedicada de 500 Mbps e IP fijas. Subtotal: $1.800.000 COP. IVA: $342.000 COP. Total: $2.142.000 COP."),
         
        ("Factura_06_Capacitacion_Inteligencia_Artificial", "Cuenta de Cobro / Factura No. CAP-304",
         "Emisor: INSTITUTO DE CIENCIA DE DATOS. Concepto: Diplomado corporativo en Retrieval-Augmented Generation y LLMs para 8 ingenieros. Subtotal: $9.500.000 COP. IVA: $1.805.000 COP. Total: $11.305.000 COP."),
         
        ("Factura_07_Renovacion_Dominios_SSL", "Factura Electrónica No. DOM-771",
         "Emisor: GO DADDY HOSTING SERVICES. Detalle: Renovación de 5 certificados SSL Wildcard y dominios corporativos .com y .co. Subtotal: $950.000 COP. IVA: $180.500 COP. Total a Pagar: $1.130.500 COP."),
         
        ("Factura_08_Mantenimiento_Aire_Acondicionado", "Factura de Servicios Generales No. CLIMA-112",
         "Emisor: CLIMATIZACIÓN INDUSTRIAL S.A.S. Concepto: Mantenimiento preventivo de unidades de precisión para Centro de Cómputo (Datacenter). Subtotal: $2.400.000 COP. IVA: $456.000 COP. Total: $2.856.000 COP."),
         
        ("Factura_09_Adquisicion_Monitores_Dell", "Factura de Hardware No. DELL-6512",
         "Emisor: TECNO DISTRIBUCIONES SANTANDER. Detalle: 10 Monitores Dell UltraSharp 27 pulgadas 4K para laboratorio de desarrollo. Subtotal: $14.500.000 COP. IVA: $2.755.000 COP. Total a Pagar: $17.255.000 COP."),
         
        ("Factura_10_Soporte_Base_Datos_Oracle", "Factura de Servicios TI No. ORA-501",
         "Emisor: DATABASE EXPERTS COLOMBIA. Objeto: Afinamiento y optimización de consultas SQL complejas para repositorio empresarial. Subtotal: $4.000.000 COP. IVA: $760.000 COP. Total: $4.760.000 COP.")
    ]

    for idx, (fname, title, body) in enumerate(facturas, 1):
        if idx % 3 == 1:
            create_pdf(FINANCE_DIR / f"{fname}.pdf", title, body)
        elif idx % 3 == 2:
            create_docx(FINANCE_DIR / f"{fname}.docx", title, body)
        else:
            create_txt(FINANCE_DIR / f"{fname}.txt", title, body)

    # ---------------- 3. TALENTO HUMANO E INFORMES (10 DOCUMENTOS) ----------------
    talento = [
        ("CV_01_Ingeniero_Software_FullStack", "Hoja de Vida: Carlos Andrés Morales",
         "Candidato: Carlos Andrés Morales | Teléfono: 315-9988776 | Correo: carlos.morales.dev@example.com\nProfesión: Ingeniero de Sistemas y Computación (UTS).\nPerfil Profesional: 4 años de experiencia en desarrollo web backend y frontend con Python, FastAPI, React, Node.js y bases de datos PostgreSQL y SQLite.\nHabilidades Clave: Python, FastAPI, Docker, Git, CI/CD, React, TypeScript, Metodologías Ágiles Scrum.\nExperiencia Laboral: Tech Solutions (2022-2025) - Líder Técnico de Microservicios."),
         
        ("CV_02_Cientifico_Datos_NLP", "Hoja de Vida: Mariana Sofía Restrepo",
         "Candidato: Mariana Sofía Restrepo | Correo: mariana.data@example.com\nProfesión: Científica de Datos y Especialista en Inteligencia Artificial.\nPerfil Profesional: Experiencia de 3 años entrenando modelos de NLP, embeddings semánticos, arquitecturas RAG y fine-tuning de Large Language Models.\nHabilidades: Python, PyTorch, HuggingFace, LangChain, ChromaDB, Pandas, Scikit-learn, SQL."),
         
        ("CV_03_Administrador_Bases_Datos_DBA", "Curriculum Vitae: Roberto Gómez Plata",
         "Candidato: Roberto Gómez Plata | Profesional Senior DBA.\nExperiencia de 6 años en afinamiento, réplicas, particionamiento de tablas y seguridad de datos relacionales y no relacionales.\nHabilidades: PostgreSQL, Oracle, SQLite, MongoDB, Redis, Linux RedHat, Backup y Disaster Recovery."),
         
        ("CV_04_Disenador_UI_UX_Figma", "Perfil Profesional: Laura Camila Duarte",
         "Candidata: Laura Camila Duarte | Diseñadora de Experiencia de Usuario.\nEspecialista en sistemas de diseño accesibles, wireframing interactivo, microinteracciones y maquetación CSS3.\nHabilidades: Figma, Adobe XD, HTML5, CSS3, Design Systems, Pruebas de Usabilidad con Usuarios."),
         
        ("CV_05_Ingeniero_DevOps_Cloud", "Hoja de Vida: David Fernando Silva",
         "Candidato: David Fernando Silva | Ingeniero DevOps y Cloud Architect.\nExperiencia de 5 años automatizando pipelines de despliegue continuo e infraestructura como código.\nHabilidades: Docker, Kubernetes, Terraform, AWS, GitHub Actions, Linux, Nginx, Monitoreo con Prometheus."),
         
        ("Informe_06_Arquitectura_Seguridad_ZeroTrust", "Informe Técnico: Evaluación de Arquitectura Zero-Trust",
         "Autor: Comité de Arquitectura Tecnológica. Resumen del Informe: Se evaluó la postura de seguridad de la infraestructura interna de la empresa. Se recomienda adoptar autenticación basada en tokens JWT con rotación obligatoria de firmas, encriptación AES-256 en reposo y políticas de mínimo privilegio."),
         
        ("Informe_07_Pruebas_Rendimiento_FastAPI", "Reporte de Pruebas de Carga y Concurrencia de API",
         "Objetivo: Evaluar el rendimiento del backend bajo carga de 500 peticiones simultáneas. Resultados: Latencia promedio de 12ms por endpoint de lectura, 0% de paquetes perdidos y consumo de memoria estable en 120MB."),
         
        ("Informe_08_Migracion_Base_Datos_Vectorial", "Especificación Técnica: Indexación Vectorial para RAG",
         "El informe describe la técnica de chunking con solapamiento (overlap) de 200 caracteres y cálculo de similitud coseno para maximizar la relevancia en búsquedas documentales no estructuradas."),
         
        ("Informe_09_Auditoria_Normativa_ISO27001", "Dictamen de Auditoría de Sistemas de Gestión de Seguridad",
         "Se verificó el cumplimiento de los controles de seguridad documental, gestión de credenciales y resguardo de variables de entorno sin exposición en repositorios públicos de Git."),
         
        ("Informe_10_Estrategia_Continuidad_Negocio_BCP", "Plan de Recuperación y Continuidad Operativa BCP/DRP",
         "Establece los tiempos de recuperación objetivo (RTO < 1 hora) y punto de recuperación objetivo (RPO < 2 horas) mediante backups automáticos diarios de la base de datos relacional y el almacén de archivos.")
    ]

    for idx, (fname, title, body) in enumerate(talento, 1):
        if idx % 3 == 1:
            create_pdf(HR_DIR / f"{fname}.pdf", title, body)
        elif idx % 3 == 2:
            create_docx(HR_DIR / f"{fname}.docx", title, body)
        else:
            create_txt(HR_DIR / f"{fname}.txt", title, body)

    print("[OK] Los 30 documentos sinteticos (.pdf, .docx, .txt) han sido generados exitosamente en test_dataset_30_docs/!")

if __name__ == "__main__":
    generate_all_30_documents()
